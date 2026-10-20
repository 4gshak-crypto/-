"""Сводит три выгрузки в данные трекера.

  python3 tools/merge_sources.py <годовой план .xls> <файл исполнителей .xlsx> <выгрузка Битрикс24 .xlsx>

Источники и что из них берётся:
  годовой план (ПТО)   — состав работ, УК, основание, оценка, факт, цвет строки (зелёный = принято, жёлтый = выполнено без акта)
  файл исполнителей    — ответственный мастер или ПТО, статус, отметка об акте, комментарии мастеров, пополнения плана
  Битрикс24            — стадия, ответственный за участок, даты, файл акта КС-2, дата создания
Плюс прежние данные трекера (tracker/data.json): вид работ, исполнитель (СВОИ/подрядчики), объекты КР.

Пишет tracker/merged.json (работы и вопросы) и tracker/merge_report.json (расхождения для разбора)."""
import collections, datetime, hashlib, json, pathlib, re, sys
import openpyxl, xlrd
sys.path.insert(0, str(pathlib.Path(__file__).parent))
from srcnorm import addr_key, work_norm, work_sim
from extract_backlog import cat_of, norm_addr, norm_uk

ROOT = pathlib.Path(__file__).resolve().parent.parent
MONTHS_RU = {'Январь': 1, 'Февраль': 2, 'Март': 3, 'Апрель': 4, 'Май': 5, 'Июнь': 6, 'Июль': 7, 'Август': 8,
             'Сентябрь': 9, 'Октябрь': 10, 'Ноябрь': 11, 'Декабрь': 12}
MASTERS = {  # фамилия -> как показывать
    'гайдуков': 'Гайдуков А.В.', 'зеленцов': 'Зеленцов А.В.', 'сандюк': 'Сандюк Ю.В.', 'вазюков': 'Вазюков В.С.',
    'шарафутдинов': 'Шарафутдинов Р.Р.', 'евграфов': 'Евграфов С.В.', 'алиев': 'Алиев Р.Р.', 'чернов': 'Чернов Р.А.',
    'наумова': 'Наумова Т.Ю.', 'братчинин': 'Братчинин Н.А.', 'мастер': 'Мастер участка (Север)',
}
sp = lambda t: re.sub(r'\s+', ' ', str(t if t is not None else '')).strip()
num = lambda v: float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) and v else None


def money(v):
    if num(v) is not None:
        return num(v)
    s = sp(v).replace('₽', '').replace(' ', '').replace('\xa0', '').replace(',', '.')
    try:
        return float(s) or None
    except ValueError:
        return None


def master_of(s):
    s = sp(s).lower()
    for k, v in MASTERS.items():
        if k in s:
            return v
    return None


# ---------- годовой план ----------
def read_plan(path):
    b = xlrd.open_workbook(path, formatting_info=True)
    s = b.sheet_by_name('2026')
    out, sec = [], None
    for r in range(s.nrows):
        a = sp(s.cell_value(r, 0))
        if a in MONTHS_RU:
            sec = a
            continue
        if not sec or a.startswith('Всего') or not (sp(s.cell_value(r, 1)) and sp(s.cell_value(r, 2))):
            continue
        xf = b.xf_list[s.cell_xf_index(r, 1)]
        col = b.colour_map.get(xf.background.pattern_colour_index)
        fill = '%02X%02X%02X' % col if col else None
        fact = sp(s.cell_value(r, 6)) if s.cell_type(r, 6) != xlrd.XL_CELL_DATE else 'дата ' + xlrd.xldate_as_datetime(s.cell_value(r, 6), b.datemode).strftime('%d.%m.%Y')
        fs = s.cell_value(r, 7)
        out.append(dict(row=r + 1, month='2026-%02d' % MONTHS_RU[sec], uk=norm_uk(s.cell_value(r, 0)), addr=norm_addr(s.cell_value(r, 1)),
                        work=sp(s.cell_value(r, 2)), basis=sp(s.cell_value(r, 3)), sum=num(s.cell_value(r, 4)),
                        fact=fact, fsum=num(fs), fsum_txt='' if num(fs) else sp(fs), act_col=sp(s.cell_value(r, 8)), fill=fill))
    return out


def plan_status(p):
    f, t, a = p['fill'], p['fact'].lower(), (p['fsum_txt'] + ' ' + p['act_col']).lower()
    if f == '008000' or f == '339966':
        return 'Принято, акт есть'
    if f == 'FF0000' or 'отмен' in t or 'отказ' in t:
        return 'Снято'
    if f == 'FFFF00' or 'выпол' in t or 'акта нет' in a or 'нет акта' in a:
        return 'Выполнено, акта нет'
    if 'в работе' in t:
        return 'В работе'
    if 'голосу' in t:
        return 'Голосование ОСС'
    return None


# ---------- файл исполнителей ----------
EX_STATUS = [('выполн', 'Выполнено, акта нет'), ('установили', 'Выполнено, акта нет'), ('в работе', 'В работе'), ('в  работе', 'В работе'),
             ('на контроле', 'На контроле'), ('зима', 'Отложено на зиму'), ('отказ', 'Снято'), ('отмена', 'Снято'),
             ('голосу', 'Голосование ОСС'), ('?', 'Требует уточнения'), ('ждем', 'В работе')]


def ex_status(g):
    g = sp(g).lower()
    if not g or re.match(r'^\d{4}-\d\d-\d\d', g):
        return None
    for k, v in EX_STATUS:
        if k in g:
            return v
    return 'В работе'  # «ремонт зонта», «альпинисты», «требуется ремонт будки» — работа идёт с оговоркой


def act_state(h):
    h = sp(h).lower()
    if not h or num(h):
        return None
    if 'акт в пто' in h:
        return 'передан в ПТО'
    if 'акт есть' in h:
        return 'есть'
    if 'электрон' in h:
        return 'электронный'
    if 'как подписыв' in h:
        return 'не знают, как подписать'
    if 'акта нет' in h or 'не сдан' in h:
        return 'нет'
    if 'по содержанию' in h:
        return 'по содержанию'
    return None


def read_exec(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    main = []
    for row in wb['Тек.Рем (08.10.2026)'].iter_rows():
        v = [c.value for c in row[:9]]
        if not (v[0] and v[1]):
            continue
        h = v[7]
        main.append(dict(src_row=row[0].row, addr=sp(v[0]), work=sp(v[1]), basis=sp(v[2]), sum=num(v[3]),
                         month=v[4].strftime('%Y-%m') if isinstance(v[4], datetime.datetime) else None,
                         who=sp(v[5]), status_raw=sp(v[6]), status=ex_status(v[6]), act=act_state(h), fsum=num(h),
                         note='' if num(h) or (act_state(h) and len(sp(h)) <= 12) else sp(h)))
    adds = []
    for name in ('пополнение от 16.10.2026', 'пополнение от 16.10.2026(доп)'):
        for r in wb[name].iter_rows(min_row=3, values_only=True):
            if not (r[2] and r[3]):
                continue
            adds.append(dict(sheet=name, uk=norm_uk(r[1]), addr=sp(r[2]), work=sp(r[3]),
                             month=r[4].strftime('%Y-%m') if isinstance(r[4], datetime.datetime) else None,
                             sum=num(r[5]), basis=sp(r[6]), master=master_of(r[7]), status_raw=sp(r[8]), status=ex_status(r[8]),
                             executor=sp(r[9]), act=sp(r[10])))
    per = []
    for name in wb.sheetnames[3:]:
        for r in wb[name].iter_rows(values_only=True):
            cells = [c for c in r if c is not None]
            if len(cells) < 4:
                continue
            txt = [sp(c) for c in cells if isinstance(c, str)]
            # адрес — первая строка вида «… ул, N», работа — следующая
            ai = next((i for i, t in enumerate(txt) if re.search(r'\d', t) and addr_key(t)[1]), None)
            if ai is None or ai + 1 >= len(txt):
                continue
            stat = next((t for t in txt if ex_status(t) and t.lower() in ('в работе', 'на контроле', 'зима', '?????', ' в работе', 'в  работе') or t.lower().strip() in ('в работе', 'на контроле', 'зима')), '')
            tail = [t for t in txt[ai + 2:] if t not in (stat,) and not master_of(t) and not t.lower().startswith('на основ')]
            per.append(dict(sheet=name, master=master_of(name), addr=txt[ai], work=txt[ai + 1], status=ex_status(stat) if stat else None,
                            status_raw=stat, note='; '.join(t for t in tail if len(t) > 3)))
    return main, adds, per


# ---------- Битрикс24 ----------
BX_STAGE = {'Успешно': 'Успешно', 'Запланировано': None, 'В работе': 'В работе', 'В работе. Утверждена дата завершения': 'В работе',
            'Просрочено': 'В работе', 'Выполнено, документы не сданы': 'Выполнено, акта нет', 'Отказано': 'Снято'}


def ddmmyyyy(s):
    m = re.match(r'(\d\d)\.(\d\d)\.(\d{4})', sp(s))
    return '%s-%s-%s' % (m.group(3), m.group(2), m.group(1)) if m else None


def read_bitrix(path):
    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    out, sec = [], None
    for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if r[0] and not r[1] and '·' in str(r[0]):
            m = re.match(r'(\S+)\s+(\d{4})', str(r[0]).strip().capitalize())
            sec = '%s-%02d' % (m.group(2), MONTHS_RU[m.group(1).capitalize()])
            continue
        if not r[1]:
            continue
        resp = sp(r[12])
        junk = (not sp(r[2]) or len(sp(r[2])) < 4 or re.search(r'тест', sp(r[2]).lower()) or re.fullmatch(r'[\d\s]+', sp(r[2]))
                or 'не выбрано' in sp(r[1]).lower() or 'тестович' in resp.lower())
        out.append(dict(src_row=i, month=sec, uk=norm_uk(r[0]), addr_raw=sp(r[1]), work=sp(r[2]), basis_type=sp(r[3]), basis_doc=sp(r[4]),
                        start=ddmmyyyy(r[5]), end=ddmmyyyy(r[6]), created=ddmmyyyy(r[7]), act_file=sp(r[8]), sum=money(r[9]), fsum=money(r[10]),
                        stage=sp(r[11]), resp=None if resp in ('', '0') or 'тестович' in resp.lower() else (master_of(resp) or resp), junk=bool(junk)))
    return out


# ---------- сопоставление ----------
def best(item, cands, used=None, min_sim=0.55):
    k = addr_key(item['addr'] if 'addr' in item else item['addr_raw'])
    best_c, best_s = None, 0
    for c in cands.get(k, []):
        if used is not None and id(c) in used:
            continue
        s = work_sim(item['work'], c['work'])
        if item.get('month') and c.get('month') and item['month'] == c['month']:
            s += 0.05
        if s > best_s:
            best_c, best_s = c, s
    return (best_c, best_s) if best_s >= min_sim else (None, best_s)


def glob_match(r, works, name, min_sim=0.9):
    """Та же работа (почти дословно) на другом адресе той же улицы или того же месяца."""
    k = addr_key(r['addr_raw'])
    best_w, best_s = None, 0
    for w in works:
        if name in w['src']:
            continue
        s = work_sim(r['work'], w['work'])
        if addr_key(w['addr'])[0] == k[0]:
            s += 0.1
        elif w['month'] != r['month']:
            continue
        if s > best_s:
            best_w, best_s = w, s
    return best_w if best_s >= min_sim + 0.1 or (best_s >= min_sim and best_w and addr_key(best_w['addr'])[0] == k[0]) else None


def index(rows, field='addr'):
    d = collections.defaultdict(list)
    for r in rows:
        d[addr_key(r[field])].append(r)
    return d


def bx_addr(a):
    """«улица Карла Либкнехта, 22» -> «Карла Либкнехта ул, 22»."""
    a = sp(a).replace('ё', 'е').replace('Ё', 'Е')
    for t, short in (('улица', 'ул'), ('проспект', 'пр-т'), ('проезд', 'пр-д'), ('шоссе', 'шоссе')):
        m = re.match(r'^(?:%s\s+(.+?)|(.+?)\s+%s),\s*(.+)$' % (t, t), a, re.I)
        if m:
            return '%s %s, %s' % (m.group(1) or m.group(2), short, m.group(3))
    return norm_addr(a)


def key_of(addr, work, month):
    h = hashlib.sha1(('%s|%s|%s' % ('|'.join(addr_key(addr)), work_norm(work), month)).encode()).hexdigest()[:10]
    return 'w' + h


def main(plan_path, exec_path, bx_path):
    plan = read_plan(plan_path)
    ex_main, ex_adds, ex_per = read_exec(exec_path)
    bx = read_bitrix(bx_path)
    prev = json.loads((ROOT / 'tracker' / 'data.json').read_text(encoding='utf-8'))
    prev_items = [dict(i, addr=i['addr'], kr=g.get('kr')) for g in prev['groups'] for i in g['items']]
    # проверенное написание адресов из прежнего трекера
    canon = {}
    for g in prev['groups']:
        canon.setdefault(addr_key(g['addr']), g['addr'])
    for b_ in json.loads((ROOT / 'tracker' / 'backlog.json').read_text(encoding='utf-8')):
        canon.setdefault(addr_key(b_['addr']), b_['addr'])
    canon.update({addr_key('50 лет ВЛКСМ д.24'): '50-летия ВЛКСМ пр-т, 24', addr_key('Высотный проезд, 2'): 'Высотный пр-д, 2'})
    rep = collections.defaultdict(list)

    works = []
    for p in plan:
        works.append(dict(p, src={'plan': p}))
    # пополнения плана, которых ещё нет в годовом плане
    idx = index([w for w in works])
    for a in ex_adds:
        c, s = best(a, idx)
        if c:
            c['src']['add'] = a
        else:
            w = dict(row=None, month=a['month'], uk=a['uk'], addr=norm_addr(a['addr']), work=a['work'], basis=a['basis'], sum=a['sum'],
                     fact='', fsum=None, fsum_txt='', act_col='', fill=None, src={'add': a})
            works.append(w)
            idx[addr_key(w['addr'])].append(w)
            rep['только в пополнении'].append(a)

    def attach(rows, name, field='addr'):
        used = set()
        for r in rows:
            c, s = best(r if field == 'addr' else dict(r, addr=r[field]), idx, used)
            if c and name not in c['src']:
                c['src'][name] = r
                used.add(id(c))
            elif c:
                rep['дубль в ' + name].append((r, c['addr'], c['work']))
            else:
                rep['не найдено: ' + name].append(r)
    attach(ex_main, 'exec')
    for r in rep.pop('не найдено: exec', []):
        new = dict(row=None, month=r['month'], uk='', addr=norm_addr(r['addr']), work=r['work'], basis=r['basis'], sum=r['sum'], fact='',
                   fsum=r['fsum'], fsum_txt='', act_col='', fill=None, src={'exec': r}, only='exec')
        works.append(new)
        idx[addr_key(new['addr'])].append(new)
        rep['только в файле исполнителей'].append(r)
    attach(ex_per, 'per')
    attach(prev_items, 'prev')
    # Битрикс: не найденное — новые работы
    used = set()
    for r in bx:
        if not r['junk'] and r['start'] and r['end'] and r['end'] < r['start']:
            rep['Битрикс24: дата завершения раньше начала'].append(r)
        if r['junk']:
            rep['Битрикс24: тестовые и пустые записи'].append(r)
            continue
        c, s = best(dict(r, addr=r['addr_raw']), idx, used)
        if c:
            c['src']['bx'] = r
            used.add(id(c))
        elif best(dict(r, addr=r['addr_raw']), idx)[0]:
            rep['Битрикс24: повторная карточка той же работы'].append(r)
        else:
            g = glob_match(r, works, 'bx')
            if g:
                g['src']['bx'] = r
                g['addr_conflict'] = r['addr_raw']
                rep['Битрикс24: другой адрес'].append((r['addr_raw'], r['work'], g['addr']))
            else:
                new = dict(row=None, month=r['month'], uk=r['uk'], addr=bx_addr(r['addr_raw']), work=r['work'],
                           basis=r['basis_type'], sum=r['sum'], fact='', fsum=r['fsum'], fsum_txt='', act_col='', fill=None, src={'bx': r}, only='bx')
                works.append(new)
                rep['только в Битрикс24'].append(r)

    uk_by_addr = {}
    for w in works:
        if w['uk']:
            uk_by_addr.setdefault(addr_key(w['addr']), w['uk'])
    out = []
    for w in works:
        S = w['src']
        p, e, a, pe, pr, b = (S.get(k) for k in ('plan', 'exec', 'add', 'per', 'prev', 'bx'))
        cand = []  # (дата источника, статус)
        if pr and pr.get('status'):
            cand.append(('06.10', {'done': 'Выполнено, акта нет'}.get(pr['status'], None) or PREV_ST.get(pr['status'])))
        if e and e['status']:
            cand.append(('08.10', e['status']))
        if pe and pe['status']:
            cand.append(('08.10 лист мастера', pe['status']))
        ps = plan_status(p) if p else None
        if ps:
            cand.append(('16.10', ps))
        if a and a['status']:
            cand.append(('16.10 пополнение', a['status']))
        bs = BX_STAGE.get(b['stage']) if b else None
        if bs == 'Успешно':
            bs = 'Принято, акт есть' if (b['act_file'] or ps == 'Принято, акт есть') else 'Выполнено, акта нет'
        if bs:
            cand.append(('19.10 Битрикс', bs))
        cand = [c for c in cand if c[1]]
        DONE = ('Выполнено, акта нет', 'Принято, акт есть')
        status = cand[-1][1] if cand else ('Не начато' if (a or (b and b['stage'] == 'Запланировано' and w.get('only') == 'bx')) else 'Нет статуса')
        conflict = ''
        strong_done = [c for c in cand if c[1] in DONE and not c[0].startswith('06.10')]
        if status not in DONE and status != 'Снято' and strong_done and cand[-1][0].startswith('19.10'):
            conflict = 'Битрикс24: «%s», а %s — %s' % (b['stage'], strong_done[-1][0], strong_done[-1][1].lower())
            status = strong_done[-1][1]
        elif status not in DONE and status != 'Снято' and strong_done:
            conflict = '%s — %s, позже %s — %s' % (strong_done[-1][0], strong_done[-1][1].lower(), cand[-1][0], cand[-1][1].lower())
        if ps == 'Принято, акт есть' and status != 'Снято':
            status = 'Принято, акт есть'  # принятие по годовому плану ПТО главнее отметок исполнителя
        if status in DONE and b and b['stage'] == 'Запланировано':
            rep['Битрикс24: выполнено, а стадия «Запланировано»'].append((w['addr'], w['work']))

        # исполнитель: СВОИ / Подрядчики
        ex = None
        if a and a['executor']:
            ex = 'Подрядчики' if 'пто' in a['executor'].lower() and 'филимон' not in a['executor'].lower() else 'СВОИ'
        if not ex and e:
            wl = e['who'].lower()
            if wl.startswith('пто'):
                ex = 'Подрядчики'
            elif 'алеев' in wl:
                ex = 'СВОИ'
        if not ex and p and p['fill'] == 'FFCC99':
            ex = 'СВОИ'
        if not ex and p and p['fill'] in ('969696', 'C0C0C0'):
            ex = 'Подрядчики'
        if not ex and pr and pr.get('ex') in ('СВОИ', 'Подрядчики'):
            ex = pr['ex']
        ex = ex or 'Не распределено'

        master = (a and a['master']) or (pe and pe['master']) or (e and master_of(e['who'])) or (b and b['resp']) or ''
        mconf = bool(b and b['resp'] and master and b['resp'] != master and b['resp'] != 'Наумова Т.Ю.')
        if b and b['resp'] == 'Наумова Т.Ю.' and master:
            rep['Битрикс24: ответственный Наумова, мастер другой'].append((w['addr'], w['work'], master))

        act = ''
        if status == 'Принято, акт есть':
            act = 'есть'
        elif e and e['act']:
            act = e['act']
        if b and b['act_file']:
            act = 'есть (КС-2 в Битрикс24)'
        notes = []
        if p and p['fact'] and p['fact'].lower() not in ('выполнено', 'выполено'):
            notes.append('План: ' + p['fact'])
        if e and e['note']:
            notes.append('Исполнитель 08.10: ' + e['note'])
        if pe and pe['note']:
            notes.append(pe['master'] + ': ' + pe['note'])
        if a and a['status_raw'] and not a['status']:
            notes.append('Пополнение: ' + a['status_raw'])
        if e and e['status_raw'] and e['status'] == 'В работе' and e['status_raw'].lower() not in ('в работе',):
            notes.append('Исполнитель: ' + e['status_raw'])
        fsum = (p and p['fsum']) or (b and b['fsum']) or (e and e['fsum']) or None
        if p and p['fsum'] and e and e['fsum'] and abs(p['fsum'] - e['fsum']) > 1:
            rep['факт различается'].append((w['addr'], w['work'], p['fsum'], e['fsum']))
        check = []
        if conflict:
            check.append('статус расходится: ' + conflict)
        if pr and pr.get('ex') in ('СВОИ', 'Подрядчики') and ex in ('СВОИ', 'Подрядчики') and pr['ex'] != ex:
            rep['исполнитель изменился'].append((w['addr'], w['work'], pr['ex'], ex, e['who'] if e else ''))
            check.append('исполнитель: по отметке 23.09 — %s, в файле исполнителей 08.10 — «%s»' % (pr['ex'], e['who'] if e else ex))
        if w.get('addr_conflict'):
            check.append('в Битрикс24 другой адрес: «%s»' % w['addr_conflict'])
        if mconf:
            check.append('мастер: у исполнителей — %s, в Битрикс24 — %s' % (master, b['resp']))
        if w.get('only') == 'exec':
            check.append('работа есть только в файле исполнителей, в годовом плане 16.10 её нет')
        if w.get('only') == 'bx':
            check.append('работа есть только в Битрикс24, в годовом плане 16.10 её нет')
        if p and p['fsum'] and e and e['fsum'] and abs(p['fsum'] - e['fsum']) > 1:
            check.append('факт: в плане %s ₽, у исполнителей %s ₽' % (('%.2f' % p['fsum']).replace('.', ','), ('%.2f' % e['fsum']).replace('.', ',')))
        fmt = lambda d: '.'.join(reversed(d.split('-'))) if d else ''
        bx_dates = ''
        if b and (b['start'] or b['end']):
            bx_dates = 'срок %s – %s' % (fmt(b['start']), fmt(b['end']))
            if b['start'] and b['end'] and b['end'] < b['start']:
                bx_dates += ' (завершение раньше начала)'
        if b and b['created']:
            bx_dates += ('; ' if bx_dates else '') + 'создана ' + fmt(b['created'])
        addr = canon.get(addr_key(w['addr'])) or w['addr']
        uk = w['uk'] or (b and b['uk']) or (pr and pr.get('uk')) or uk_by_addr.get(addr_key(w['addr']), '')
        item = dict(key=key_of(w['addr'], w['work'], w['month']), row=w['row'], month=w['month'], addr=addr, uk=uk,
                    work=w['work'], cat=(pr and pr.get('cat')) or cat_of(w['work']), basis=w['basis'] or (b and b['basis_type']) or '',
                    sum=w['sum'] or (b and b['sum']) or None, fsum=fsum, ex=ex, master=master, status=status, act=act,
                    kr='; '.join(pr['kr']['works']) if pr and pr.get('kr') else '',
                    bx_stage=b['stage'] if b else '', bx_start=b['start'] if b else '', bx_end=b['end'] if b else '', bx_created=b['created'] if b else '',
                    in_bx=bool(b), in_exec=bool(e or a or pe), conflict=conflict, history=' → '.join('%s: %s' % c for c in cand),
                    src_note=' · '.join(([item_hist] if (item_hist := ' → '.join('%s: %s' % c for c in cand)) else []) + notes),
                    master_conflict=b['resp'] if mconf else '', check='; '.join(check), bx_dates=bx_dates,
                    addr_conflict=w.get('addr_conflict', ''), only=w.get('only', ''))
        out.append(item)
        if conflict:
            rep['статусы расходятся'].append((w['addr'], w['work'], conflict))
        if mconf:
            rep['мастер расходится'].append((w['addr'], w['work'], master, b['resp']))
    keys = collections.Counter(i['key'] for i in out)
    for i in out:
        if keys[i['key']] > 1:  # одинаковые работы по одному адресу в одном месяце
            i['key'] += '-%s' % (i['row'] or len(out))
    dst = ROOT / 'tracker'
    (dst / 'merged.json').write_text(json.dumps({'generated': datetime.date.today().isoformat(),
                                                 'sources': {'plan': '16.10.2026', 'exec': '08.10.2026 + пополнения 16.10', 'bitrix': '19.10.2026'},
                                                 'items': out}, ensure_ascii=False, indent=1), encoding='utf-8')
    (dst / 'merge_report.json').write_text(json.dumps(rep, ensure_ascii=False, indent=1, default=str), encoding='utf-8')
    print(len(out), 'работ;', {k: len(v) for k, v in rep.items()})


PREV_ST = {'blocked': 'Нужно решение', 'unclear': 'Требует уточнения', 'nostatus': None, 'capital': 'Вопрос капремонта', 'vote': 'Голосование ОСС',
           'todo': 'Не начато', 'assigned': 'Передано исполнителю', 'watch': 'На контроле', 'material': 'Материал заказан', 'work': 'В работе',
           'winter': 'Отложено на зиму', 'cancel': 'Снято'}

if __name__ == '__main__':
    main(*sys.argv[1:4])
