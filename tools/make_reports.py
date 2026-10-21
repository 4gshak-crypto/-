"""Сводный трекер и перечень вопросов в Excel.
Читает tracker/merged.json и tracker/questions.json, пишет в reports/:
  Сводный трекер текущего ремонта <дата>.xlsx
  Вопросы для трекера <дата>.xlsx"""
import calendar, datetime, json, pathlib
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = pathlib.Path(__file__).resolve().parent.parent
M = json.loads((ROOT / 'tracker' / 'merged.json').read_text(encoding='utf-8'))
Q = json.loads((ROOT / 'tracker' / 'questions.json').read_text(encoding='utf-8'))
DATE = datetime.date.fromisoformat(Q['generated'])
OUT = ROOT / 'reports'
OUT.mkdir(exist_ok=True)

F = 'Arial'
H_FILL = PatternFill('solid', fgColor='1F3A4D')
H_FONT = Font(name=F, bold=True, color='FFFFFF', size=10)
BASE = Font(name=F, size=10)
BOLD = Font(name=F, size=10, bold=True)
TITLE = Font(name=F, size=14, bold=True)
SECTION = Font(name=F, size=12, bold=True)
NOTE = Font(name=F, size=9, italic=True, color='5B665F')
INPUT = PatternFill('solid', fgColor='FFF7D6')
thin = Side(style='thin', color='D9DDD6')
GRID = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical='top')
# стадии — те же, что на странице трекера: стадия, статусы трекера, цвет
STAGES = [
    ('Не начато', ['Не начато', 'Передано исполнителю'], 'ECEEF0'),
    ('В работе', ['В работе', 'На контроле', 'Материал заказан'], 'E1EBF7'),
    ('Выполнено, ждёт акта', ['Выполнено, акта нет'], 'FAF0D2'),
    ('Принято с актом', ['Принято, акт есть'], 'E0F1E2'),
    ('Отложено на зиму', ['Отложено на зиму'], 'EFE8F8'),
    ('Нужно решение', ['Нужно решение', 'Требует уточнения', 'Нет статуса', 'Вопрос капремонта', 'Голосование ОСС'], 'FBE5E1'),
    ('Снято с плана', ['Снято'], 'ECEEEC'),
]
STAGE_OF = {s: st for st, sts, _ in STAGES for s in sts}
COLOR = {st: c for st, _, c in STAGES}
PLAN = [st for st, _, _ in STAGES if st != 'Снято с плана']  # шесть стадий, из которых состоит план
SHORT = ['Не начато', 'В работе', 'Ждёт акта', 'Принято', 'На зиму', 'Нужно решение']
OPEN_ST = ['Не начато', 'В работе', 'Нужно решение']  # открытые стадии: только они бывают просрочены
MN = ['Январь', 'Февраль', 'Март', 'Апрель', 'Май', 'Июнь', 'Июль', 'Август', 'Сентябрь', 'Октябрь', 'Ноябрь', 'Декабрь']


def header(ws, row, names, widths=None):
    for c, n in enumerate(names, 1):
        cell = ws.cell(row=row, column=c, value=n)
        cell.font, cell.fill, cell.alignment, cell.border = H_FONT, H_FILL, Alignment(wrap_text=True, vertical='center'), GRID
        if widths:
            ws.column_dimensions[get_column_letter(c)].width = widths[c - 1]
    ws.row_dimensions[row].height = 30


def body(cell, fmt=None, bold=False):
    cell.font = BOLD if bold else BASE
    cell.alignment, cell.border = WRAP, GRID
    if fmt:
        cell.number_format = fmt


def fill(color):  # заливка для условного форматирования: Excel берёт цвет фона
    return PatternFill(start_color=color, end_color=color, fill_type='solid')


def printable(book):
    for sh in book.worksheets:
        sh.page_setup.orientation = 'landscape'
        sh.page_setup.fitToWidth, sh.page_setup.fitToHeight = 1, 0
        sh.sheet_properties.pageSetUpPr.fitToPage = True


def month_end(m):
    y, mm = map(int, m.split('-'))
    return datetime.date(y, mm, calendar.monthrange(y, mm)[1])


def stage(i):
    return STAGE_OF.get(i['status'], 'Нужно решение')


# ======================= сводный трекер =======================
wb = Workbook()
sv = wb.active
sv.title = 'Сводка'
ws = wb.create_sheet('Работы')
cols = ['№', 'Месяц плана', 'Срок (конец месяца)', 'Адрес', 'Управляющая компания', 'Работа', 'Вид работ', 'Основание', 'Оценка, ₽', 'Факт, ₽', 'Исполнитель',
        'Мастер участка', 'Стадия', 'Статус', 'Акт', 'Просрочено, дней', 'Стадия Битрикс24', 'Сроки Битрикс24', 'Сверить', 'Источники статуса',
        'Строка годового плана']
widths = [5, 11, 11, 26, 13, 44, 18, 30, 12, 12, 13, 20, 20, 20, 18, 11, 16, 26, 40, 50, 9]
L = {n: get_column_letter(k) for k, n in enumerate(cols, 1)}
assert (L['Стадия'], L['Статус'], L['Просрочено, дней'], L['Стадия Битрикс24'], L['Сверить']) == ('M', 'N', 'P', 'Q', 'S')
header(ws, 1, cols, widths)
NST = sum(len(s[1]) for s in STAGES)  # строк в таблице «статус → стадия» на листе «Словарь»
items = sorted(M['items'], key=lambda i: (i['month'], i['addr'], i['row'] or 9999))
for r, i in enumerate(items, 2):
    vals = [r - 1, MN[int(i['month'][5:]) - 1] + ' ' + i['month'][:4], month_end(i['month']), i['addr'], i['uk'], i['work'], i['cat'], i['basis'],
            i['sum'], i['fsum'], i['ex'], i['master'] or '(не назначен)', None, i['status'], i['act'], None, i['bx_stage'], i['bx_dates'], i['check'],
            i['src_note'], i['row']]
    for c, v in enumerate(vals, 1):
        body(ws.cell(row=r, column=c, value=v), '#,##0' if c in (9, 10) else 'DD.MM.YYYY' if c == 3 else None, bold=(c == 13))
    # стадия — по статусу из «Словаря»; просрочка — от сегодняшней даты открытия файла
    ws['M%d' % r] = '=IFERROR(VLOOKUP(N{0},Словарь!$A$2:$B${1},2,FALSE),"Нужно решение")'.format(r, NST + 1)
    ws['P%d' % r] = '=IF(AND(OR(M{0}="Не начато",M{0}="В работе",M{0}="Нужно решение"),C{0}<TODAY()),TODAY()-C{0},"")'.format(r)
    ws['P%d' % r].number_format = '0'
last = len(items) + 1
ws.freeze_panes = 'E2'
ws.auto_filter.ref = 'A1:%s%d' % (get_column_letter(len(cols)), last)
for st in PLAN[1:] + ['Снято с плана']:  # «Не начато» без заливки
    ws.conditional_formatting.add('A2:%s%d' % (get_column_letter(len(cols)), last), FormulaRule(formula=['$M2="%s"' % st], fill=fill(COLOR[st])))
ws.conditional_formatting.add('P2:P%d' % last, FormulaRule(formula=['AND(ISNUMBER($P2),$P2>0)'], font=Font(name=F, bold=True, color='C0392B')))

# словарь: статус → стадия (на него ссылаются формулы) и соответствие статусов источников
sl = wb.create_sheet('Словарь')
header(sl, 1, ['Статус в трекере', 'Стадия (как на странице трекера)', 'Может быть просрочена'])
r = 2
for st, sts, _ in STAGES:
    for s in sts:
        for c, v in enumerate([s, st, 'да' if st in OPEN_ST else ''], 1):
            body(sl.cell(row=r, column=c, value=v))
        r += 1
assert r == NST + 2
sl.cell(row=r + 1, column=1, value='Соответствие статусов в источниках и в трекере').font = BOLD
header(sl, r + 2, ['Источник', 'Как записано в источнике', 'Статус в трекере', 'Примечание'], [40, 34, 24, 50])
MAP = [
    ('Годовой план ПТО', 'зелёная строка', 'Принято, акт есть', ''),
    ('Годовой план ПТО', 'жёлтая строка, «Выполнено», «Акта нет»', 'Выполнено, акта нет', ''),
    ('Годовой план ПТО', 'красная строка, «отмена», «отказ»', 'Снято', ''),
    ('Годовой план ПТО', 'оранжевая строка', '(исполнитель УК)', 'цвет говорит об исполнителе, а не о статусе'),
    ('Годовой план ПТО', 'серая строка', '(исполнитель — Подрядчики)', 'цвет говорит об исполнителе, а не о статусе'),
    ('Файл исполнителей', 'в столбце исполнителя «ПТО»', '(исполнитель — Подрядчики)', ''),
    ('Файл исполнителей', 'в столбце исполнителя мастер или электрик', '(исполнитель УК)', 'работа своими силами'),
    ('Файл исполнителей', 'выполнено', 'Выполнено, акта нет', 'пока ПТО не подтвердит акт'),
    ('Файл исполнителей', 'в работе', 'В работе', ''),
    ('Файл исполнителей', 'на контроле', 'На контроле', 'предложение: заменить на «В работе» с комментарием, что мешает'),
    ('Файл исполнителей', 'зима', 'Отложено на зиму', ''),
    ('Файл исполнителей', 'отказ, отмена', 'Снято', ''),
    ('Файл исполнителей', '?????', 'Требует уточнения', ''),
    ('Файл исполнителей', 'Акт в ПТО / Акта нет / как подписывать', '(поле «Акт»)', 'это состояние акта, а не статус работы'),
    ('Битрикс24', 'Запланировано', 'не меняет статус', 'стоит и у выполненных работ — стадии не ведутся'),
    ('Битрикс24', 'В работе, Просрочено, В работе. Утверждена дата', 'В работе', ''),
    ('Битрикс24', 'Выполнено, документы не сданы', 'Выполнено, акта нет', ''),
    ('Битрикс24', 'Успешно', 'Принято, акт есть — если есть КС-2 или зелёная строка, иначе «Выполнено, акта нет»', 'смысл стадии нужно подтвердить'),
    ('Битрикс24', 'Отказано', 'Снято', ''),
]
for k, row in enumerate(MAP, r + 3):
    for c, v in enumerate(row, 1):
        body(sl.cell(row=k, column=c, value=v))
sl.cell(row=k + 2, column=1, value='Предложение по соответствию — для утверждения (вопрос о едином словаре статусов).').font = NOTE

# --- сводка формулами по листу «Работы»: те же блоки, что на странице трекера
W = "Работы!"
rng = lambda col: "%s$%s$2:$%s$%d" % (W, col, col, last)
ST, OD, DUE = rng('M'), rng('P'), rng('C')
open_of = lambda col, crit: '+'.join('COUNTIFS(%s,%s,%s,"%s")' % (rng(col), crit, ST, s) for s in OPEN_ST)


def row_out(n, spans, values, hdr=False, fmts=None, bold=(), height=None):
    """Строка сводки: spans — столбцы (первый, последний) для каждого значения, широкие ячейки объединяются."""
    for k, ((c1, c2), v) in enumerate(zip(spans, values)):
        sv.cell(row=n, column=c1, value=v)
        for c in range(c1, c2 + 1):
            x = sv.cell(row=n, column=c)
            if hdr:
                x.font, x.fill, x.alignment, x.border = H_FONT, H_FILL, Alignment(wrap_text=True, vertical='center'), GRID
            else:
                body(x, fmts[k] if fmts else None, bold=k in bold)
        if c2 > c1:
            sv.merge_cells(start_row=n, start_column=c1, end_row=n, end_column=c2)
    if height:
        sv.row_dimensions[n].height = height


def section(n, title, sub=None):
    sv.cell(row=n, column=1, value=title).font = SECTION
    if sub:
        sv.cell(row=n + 1, column=1, value=sub).font = NOTE
    return n + (3 if sub else 2)


sv['A1'] = 'Сводный трекер текущего ремонта'
sv['A1'].font = TITLE
sv['A2'] = ('Сведено %s из трёх выгрузок: годовой план ПТО от 16.10.2026, файл исполнителей от 08.10.2026 с пополнениями от 16.10.2026, '
            'Битрикс24 от 19.10.2026. Отметки прежних снимков (23.09, 06.10) не используются. Все цифры считаются формулами по листу «Работы», '
            'просрочка — на сегодняшнюю дату. Стадии и проблемы — те же, что на странице трекера.') % DATE.strftime('%d.%m.%Y')
sv['A2'].font, sv['A2'].alignment = NOTE, WRAP
sv.merge_cells('A2:L2')
sv.row_dimensions[2].height = 42

# 1. где мы сейчас
r = section(4, 'Где мы сейчас — работы плана по стадиям', 'Стадии идут по ходу работы: не начато → в работе → выполнено, ждёт акта → принято с актом. '
            '«Отложено на зиму» и «Нужно решение» — вне хода. Снятые с плана в проценты не входят.')
S1 = [(1, 1), (2, 2), (3, 3), (4, 4), (5, 8), (9, 12)]
row_out(r, S1, ['Стадия', 'Работ', '% плана', 'Из них с проблемой', 'Какая проблема', 'Что значит стадия'], hdr=True)
STAGE_PROB = {
    'Не начато': ('=COUNTIFS(%s,"Не начато",%s,">0")' % (ST, OD), 'месяц плана прошёл'),
    'В работе': ('=COUNTIFS(%s,"В работе",%s,">0")' % (ST, OD), 'дольше плана: месяц плана прошёл'),
    'Выполнено, ждёт акта': ('=COUNTIFS(%s,"Выполнено, ждёт акта",%s,"<"&TODAY())' % (ST, DUE), 'акт задерживается: месяц плана прошёл'),
    'Нужно решение': ('=COUNTIFS(%s,"Нужно решение",%s,">0")' % (ST, OD), 'месяц плана прошёл'),
}
MEANING = {
    'Не начато': 'Работа ещё не начата: «Не начато», «Передано исполнителю»',
    'В работе': 'Идёт работа: «В работе», «На контроле», «Материал заказан»',
    'Выполнено, ждёт акта': 'Исполнитель сделал работу, акт ещё не подписан',
    'Принято с актом': 'Работа закрыта актом',
    'Отложено на зиму': 'Перенесено на тёплый сезон',
    'Нужно решение': 'Работа стоит: «Нужно решение», «Требует уточнения», «Нет статуса», «Голосование ОСС»',
}
top, tot = r + 1, r + 1 + len(PLAN)  # первая стадия и строка «Итого в плане»
for k, st in enumerate(PLAN):
    n = top + k
    pf, pt = STAGE_PROB.get(st, (None, '—'))
    row_out(n, S1, [st, '=COUNTIF(%s,"%s")' % (ST, st), '=IF($B$%d=0,0,B%d/$B$%d)' % (tot, n, tot), pf, pt, MEANING[st]],
            fmts=[None, '#,##0', '0%', '#,##0', None, None], bold=(0, 1), height=28)
    sv.cell(row=n, column=1).fill = PatternFill('solid', fgColor=COLOR[st])
row_out(tot, S1, ['Итого в плане', '=SUM(B%d:B%d)' % (top, tot - 1), '=IF(B%d=0,0,1)' % tot, '=COUNTIF(%s,">0")' % OD,
                  'просрочено всего', 'Месяц плана прошёл, а работа не начата, в работе или стоит'],
        fmts=[None, '#,##0', '0%', '#,##0', None, None], bold=(0, 1, 3), height=28)
row_out(tot + 1, S1, ['Снято с плана', '=COUNTIF(%s,"Снято с плана")' % ST, None, None, None, 'В проценты не входит'],
        fmts=[None, '#,##0', None, None, None, None])
# число проблемных — красным, если не ноль
sv.conditional_formatting.add('D%d:D%d' % (top, tot), CellIsRule(operator='greaterThan', formula=['0'], font=Font(name=F, bold=True, color='C0392B')))
r = tot + 3
row_out(r, [(1, 1), (2, 2), (3, 3)], ['Итоги', 'Значение', '% плана'], hdr=True)
done_n, acc_n = top + PLAN.index('Выполнено, ждёт акта'), top + PLAN.index('Принято с актом')
KEY = [
    ('Выполнено: ждёт акта + принято с актом', '=B%d+B%d' % (done_n, acc_n), True),
    ('Принято с актом', '=B%d' % acc_n, True),
    ('Оценка плана, ₽', '=SUMIFS(%s,%s,"<>Снято с плана")' % (rng('I'), ST), False),
    ('Факт, ₽', '=SUMIFS(%s,%s,"<>Снято с плана")' % (rng('J'), ST), False),
    ('Работ без оценки', '=B%d-COUNTIFS(%s,"<>Снято с плана",%s,">0")' % (tot, ST, rng('I')), False),
]
for k, (label, f, pct) in enumerate(KEY, r + 1):
    row_out(k, [(1, 1), (2, 2), (3, 3)], [label, f, '=IF($B$%d=0,0,B%d/$B$%d)' % (tot, k, tot) if pct else None],
            fmts=[None, '#,##0', '0%'], bold=(1, 2))

# 2. где проблемы
r = section(k + 2, 'Где проблемы — по важности', 'Красное срывает сроки, жёлтое копится, серое — порядок в данных. '
            'Сами работы — фильтром на листе «Работы» (столбец справа подсказывает, по каким столбцам).')
S2 = [(1, 1), (2, 2), (3, 4), (5, 8), (9, 12)]
row_out(r, S2, ['Проблема', 'Работ', 'Важность', 'Кто решает', 'Как найти на листе «Работы»'], hdr=True)
SEV = {'high': ('срывает сроки', 'FBE5E1'), 'mid': ('копится', 'FAF0D2'), 'low': ('порядок в данных', 'ECEEF0')}
PROBS = [
    ('Не начаты, хотя месяц плана прошёл', '=COUNTIFS(%s,"Не начато",%s,">0")' % (ST, OD), 'high',
     'Исполнитель и ПТО: начать или перенести срок', 'Стадия «Не начато», «Просрочено, дней» больше 0'),
    ('В работе дольше плана', '=COUNTIFS(%s,"В работе",%s,">0")' % (ST, OD), 'high',
     'Мастер: дата завершения и что мешает', 'Стадия «В работе», «Просрочено, дней» больше 0'),
    ('Работа стоит — нужно решение', '=COUNTIF(%s,"Нужно решение")' % ST, 'high',
     'Руководство и ПТО: решить или снять с плана', 'Стадия «Нужно решение»'),
    ('Выполнены, но нет акта с прошлых месяцев', '=COUNTIFS(%s,"Выполнено, ждёт акта",%s,"<"&TODAY())' % (ST, DUE), 'mid',
     'ПТО: оформить и подписать акты', 'Стадия «Выполнено, ждёт акта», месяц плана прошёл'),
    ('Открытые работы без мастера участка', '=' + open_of('L', '"(не назначен)"'), 'mid',
     'Назначить ответственного мастера', 'Мастер «(не назначен)», стадия не начато, в работе или нужно решение'),
    ('Открытые работы без исполнителя', '=' + open_of('K', '"Не распределено"'), 'mid',
     'ПТО: решить, УК или подрядчик', 'Исполнитель «Не распределено», стадия не начато, в работе или нужно решение'),
    ('Расходятся план, файл исполнителей и Битрикс24', '=COUNTIF(%s,"?*")' % rng('S'), 'low',
     'Администратор: сверить источники', 'Столбец «Сверить» не пустой; список — на листе «Сверка»'),
    ('Выполнено, а в Битрикс24 стадия «Запланировано»', '=COUNTIFS({0},"Выполнено, ждёт акта",{1},"Запланировано")+COUNTIFS({0},"Принято с актом",{1},"Запланировано")'
     .format(ST, rng('Q')), 'low', 'Ответственный за Битрикс24: обновить стадии', '«Стадия Битрикс24» — «Запланировано», стадия выполнено или принято'),
]
for k, (label, f, sev, who, how) in enumerate(PROBS, r + 1):
    row_out(k, S2, [label, f, SEV[sev][0], who, how], fmts=[None, '#,##0', None, None, None], bold=(0, 1), height=28)
    for c in (3, 4):
        sv.cell(row=k, column=c).fill = PatternFill('solid', fgColor=SEV[sev][1])


# 3. по месяцам и кто отстаёт
def overdue(i):
    return stage(i) in OPEN_ST and month_end(i['month']) < DATE


def laggards(key):  # сначала — у кого больше просроченных, затем меньший % выполнения
    g = {}
    for i in M['items']:
        if stage(i) == 'Снято с плана':
            continue
        x = g.setdefault(key(i), [0, 0, 0])
        x[0] += 1
        x[1] += stage(i) in ('Выполнено, ждёт акта', 'Принято с актом')
        x[2] += overdue(i)
    return sorted(g, key=lambda k: (-g[k][2], g[k][1] / g[k][0], k))


def table(start, title, sub, first, col, groups, months=False):
    hrow = section(start, title, sub)
    header(sv, hrow, [first, 'В плане'] + SHORT + ['Просрочено', '% выполнено', '% принято с актом'] + (['Состояние месяца'] if months else []))
    first_row, last_row = hrow + 1, hrow + len(groups)
    for n, (label, m) in enumerate(groups, first_row):
        crit = '"%s"' % label
        f = ['=COUNTIFS(%s,%s,%s,"<>Снято с плана")' % (rng(col), crit, ST)]
        f += ['=COUNTIFS(%s,%s,%s,"%s")' % (rng(col), crit, ST, st) for st in PLAN]
        f += ['=COUNTIFS(%s,%s,%s,">0")' % (rng(col), crit, OD),
              '=IF(B{0}=0,"",(E{0}+F{0})/B{0})'.format(n), '=IF(B{0}=0,"",F{0}/B{0})'.format(n)]
        if months:  # как на странице: просрочен / без акта N / закрыт / идёт сейчас / впереди
            e = month_end(m)
            f.append('=IF(DATE({1},{2},{3})<TODAY(),IF(C{0}+D{0}+H{0}>0,"просрочен",IF(E{0}>0,"без акта "&E{0},"закрыт")),'
                     'IF(DATE({1},{2},1)<=TODAY(),"идёт сейчас","впереди"))'.format(n, e.year, e.month, e.day))
        body(sv.cell(row=n, column=1, value=label), bold=True)
        for c, ff in enumerate(f, 2):
            body(sv.cell(row=n, column=c, value=ff), '0%' if c in (10, 11) else '#,##0', bold=(c in (9, 10, 12)))
    sv.conditional_formatting.add('B%d:I%d' % (first_row, last_row), CellIsRule(operator='equal', formula=['0'], font=Font(name=F, color='A3A8A4')))
    sv.conditional_formatting.add('I%d:I%d' % (first_row, last_row), CellIsRule(operator='greaterThan', formula=['0'], font=Font(name=F, bold=True, color='C0392B')))
    sv.conditional_formatting.add('J%d:J%d' % (first_row, last_row), DataBarRule(start_type='num', start_value=0, end_type='num', end_value=1, color='2E7D3A'))
    if months:
        for test, color, ink in [('$L{0}="просрочен"', 'FBE5E1', 'A8321F'), ('LEFT($L{0},8)="без акта"', 'FAF0D2', '7F5B00'),
                                 ('$L{0}="закрыт"', 'E0F1E2', '256B30'), ('$L{0}="идёт сейчас"', 'E1EBF7', '164F93'), ('$L{0}="впереди"', 'ECEEF0', '4F5860')]:
            sv.conditional_formatting.add('L%d:L%d' % (first_row, last_row),
                                          FormulaRule(formula=[test.format(first_row)], fill=fill(color), font=Font(name=F, bold=True, color=ink)))
    return last_row + 2


months = sorted({i['month'] for i in M['items']})
r = table(k + 2, 'По месяцам плана', '«Просрочен» — месяц закончился, а работы не сделаны. «Без акта» — всё сделано, но не все акты оформлены. '
          '«Закрыт» — всё принято с актом.', 'Месяц плана', 'B', [(MN[int(m[5:]) - 1] + ' ' + m[:4], m) for m in months], months=True)
r = table(r, 'Кто отстаёт: исполнители', 'Сначала — у кого больше просроченных работ (на дату сведения).', 'Исполнитель', 'K',
          [(x, None) for x in laggards(lambda i: i['ex'])])
r = table(r, 'Кто отстаёт: мастера участков', None, 'Мастер участка', 'L', [(x, None) for x in laggards(lambda i: i['master'] or '(не назначен)')])
r = table(r, 'Кто отстаёт: управляющие компании', None, 'Управляющая компания', 'E', [(x, None) for x in laggards(lambda i: i['uk'])])
sv.cell(row=r, column=1, value='«В плане» — без снятых работ. «Просрочено» — работа не начата, в работе или стоит, а месяц плана закончился; '
        'считается от сегодняшней даты. «% выполнено» = (ждёт акта + принято с актом) / в плане.').font = NOTE
sv.column_dimensions['A'].width = 44
for c in 'BCDEFGH':
    sv.column_dimensions[c].width = 11
for c in 'IJK':
    sv.column_dimensions[c].width = 13
sv.column_dimensions['L'].width = 15

# --- лист сверки
sc = wb.create_sheet('Сверка')
header(sc, 1, ['Адрес', 'Работа', 'Месяц', 'Статус в трекере', 'Мастер', 'Что расходится', 'Решение'], [26, 44, 12, 20, 20, 60, 30])
chk = [i for i in items if i['check']]
for n, i in enumerate(chk, 2):
    for c, v in enumerate([i['addr'], i['work'], MN[int(i['month'][5:]) - 1], i['status'], i['master'], i['check'], None], 1):
        body(sc.cell(row=n, column=c, value=v))
    sc.cell(row=n, column=7).fill = INPUT
sc.freeze_panes = 'C2'
sc.auto_filter.ref = 'A1:G%d' % (len(chk) + 1)
sc.cell(row=len(chk) + 3, column=1, value='Столбец «Решение» (жёлтый) заполняет ПТО или администратор трекера, например: «верно по плану», «адрес — Пушкарева 74».').font = NOTE
wb.move_sheet('Словарь', offset=2)
name1 = OUT / ('Сводный трекер текущего ремонта %s.xlsx' % DATE.strftime('%d.%m.%Y'))
printable(wb)
wb.save(name1)


# ======================= вопросы =======================
wq = Workbook()
qs = wq.active
qs.title = 'Вопросы'
qs['A1'] = 'Вопросы, без ответа на которые трекер не показывает ход работ'
qs['A1'].font = TITLE
qs['A2'] = ('Составлено %s по сверке годового плана (16.10), файла исполнителей (08.10, пополнения 16.10) и выгрузки Битрикс24 (19.10). '
            'Ответы пишите в жёлтые столбцы: «Ответ» и «Кто ответил». Например: «Акт № 15 от 07.10.2026», «перенести на ноябрь», «снять с плана». '
            'Список работ по каждому вопросу — на листе «Перечни».') % DATE.strftime('%d.%m.%Y')
qs['A2'].font, qs['A2'].alignment = NOTE, WRAP
qs.merge_cells('A2:I2')
qs.row_dimensions[2].height = 44
header(qs, 4, ['№', 'Кому', 'Тема', 'Вопрос', 'Зачем это трекеру', 'Работ в перечне', 'Срок ответа', 'Ответ', 'Кто ответил'],
       [5, 20, 14, 70, 48, 11, 13, 40, 18])
for n, q in enumerate(Q['questions'], 5):
    vals = [q['n'], q['to'], q['topic'], q['text'], q['why'], None, q['term'], None, None]
    for c, v in enumerate(vals, 1):
        body(qs.cell(row=n, column=c, value=v), bold=(c == 1))
    # число работ считается по листу «Перечни»
    qs.cell(row=n, column=6, value='=IF(COUNTIF(Перечни!$A:$A,A%d)=0,"—",COUNTIF(Перечни!$A:$A,A%d))' % (n, n))
    for c in (8, 9):
        qs.cell(row=n, column=c).fill = INPUT
qs.freeze_panes = 'D5'
qs.auto_filter.ref = 'A4:I%d' % (4 + len(Q['questions']))

pe = wq.create_sheet('Перечни')
header(pe, 1, ['№ вопроса', 'Адрес', 'Работа', 'Месяц', 'Статус в трекере', 'Мастер', 'Исполнитель', 'Что уточнить', 'Ответ'],
       [9, 26, 44, 11, 20, 20, 13, 46, 30])
n = 2
for q in Q['questions']:
    for it in q['items']:
        for c, v in enumerate([q['n'], it['addr'], it['work'], it['month'], it['status'], it['master'], it['ex'], it['note'], None], 1):
            body(pe.cell(row=n, column=c, value=v))
        pe.cell(row=n, column=9).fill = INPUT
        n += 1
pe.freeze_panes = 'C2'
pe.auto_filter.ref = 'A1:I%d' % (n - 1)

# словарь статусов — на него ссылается вопрос о едином словаре
sd = wq.create_sheet('Словарь статусов')
header(sd, 1, ['Источник', 'Как записано в источнике', 'Статус в трекере', 'Примечание'], [22, 40, 34, 50])
for r, row in enumerate(MAP, 2):
    for c, v in enumerate(row, 1):
        body(sd.cell(row=r, column=c, value=v))
sd.cell(row=len(MAP) + 3, column=1, value='Предложение для утверждения. Если согласны — напишите «согласовано» в ответ на вопрос о словаре статусов.').font = NOTE
name2 = OUT / ('Вопросы для трекера %s.xlsx' % DATE.strftime('%d.%m.%Y'))
printable(wq)
wq.save(name2)
print(name1)
print(name2)
