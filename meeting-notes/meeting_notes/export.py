"""Экспорт протокола: Markdown и Word (DOCX)."""

from __future__ import annotations

from pathlib import Path

from .summarize import MeetingSummary
from .transcribe import Transcript


def summary_to_markdown(summary: MeetingSummary, meeting_date: str = "", source_name: str = "") -> str:
    lines: list[str] = [f"# {summary.title or 'Протокол совещания'}", ""]
    meta = []
    if meeting_date:
        meta.append(f"**Дата:** {meeting_date}")
    if summary.participants:
        meta.append("**Участники:** " + ", ".join(summary.participants))
    if source_name:
        meta.append(f"**Запись:** {source_name}")
    if meta:
        lines += meta + [""]

    lines += ["## Кратко", "", summary.summary or "—", ""]

    if summary.topics:
        lines += ["## Обсуждали", ""]
        for i, t in enumerate(summary.topics, 1):
            lines.append(f"### {i}. {t.title}")
            lines.append("")
            lines.append(t.discussion)
            lines.append("")

    lines += ["## Решения", ""]
    lines += [f"{i}. {d}" for i, d in enumerate(summary.decisions, 1)] or ["—"]
    lines.append("")

    lines += ["## Поручения", ""]
    if summary.action_items:
        lines += ["| № | Задача | Ответственный | Срок |", "|---|---|---|---|"]
        for i, a in enumerate(summary.action_items, 1):
            lines.append(f"| {i} | {_cell(a.task)} | {_cell(a.owner)} | {_cell(a.due)} |")
    else:
        lines.append("—")
    lines.append("")

    lines += ["## Открытые вопросы", ""]
    lines += [f"- {q}" for q in summary.open_questions] or ["—"]
    lines.append("")

    if summary.next_steps:
        lines += ["## Дальше", "", summary.next_steps, ""]
    return "\n".join(lines)


def _cell(text: str) -> str:
    return text.replace("|", "／").replace("\n", " ").strip()


def write_docx(
    path: Path,
    summary: MeetingSummary,
    transcript: Transcript | None = None,
    meeting_date: str = "",
) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    heading = doc.add_heading(summary.title or "Протокол совещания", level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT

    if meeting_date:
        p = doc.add_paragraph()
        p.add_run("Дата: ").bold = True
        p.add_run(meeting_date)
    if summary.participants:
        p = doc.add_paragraph()
        p.add_run("Участники: ").bold = True
        p.add_run(", ".join(summary.participants))

    doc.add_heading("Кратко", level=1)
    doc.add_paragraph(summary.summary or "—")

    if summary.topics:
        doc.add_heading("Обсуждали", level=1)
        for i, t in enumerate(summary.topics, 1):
            doc.add_heading(f"{i}. {t.title}", level=2)
            doc.add_paragraph(t.discussion)

    doc.add_heading("Решения", level=1)
    if summary.decisions:
        for d in summary.decisions:
            doc.add_paragraph(d, style="List Number")
    else:
        doc.add_paragraph("—")

    doc.add_heading("Поручения", level=1)
    if summary.action_items:
        table = doc.add_table(rows=1, cols=4)
        table.style = "Light Grid Accent 1"
        for cell, text in zip(table.rows[0].cells, ("№", "Задача", "Ответственный", "Срок")):
            cell.text = text
            for run in cell.paragraphs[0].runs:
                run.bold = True
        for i, a in enumerate(summary.action_items, 1):
            row = table.add_row().cells
            row[0].text = str(i)
            row[1].text = a.task
            row[2].text = a.owner
            row[3].text = a.due
    else:
        doc.add_paragraph("—")

    doc.add_heading("Открытые вопросы", level=1)
    if summary.open_questions:
        for q in summary.open_questions:
            doc.add_paragraph(q, style="List Bullet")
    else:
        doc.add_paragraph("—")

    if summary.next_steps:
        doc.add_heading("Дальше", level=1)
        doc.add_paragraph(summary.next_steps)

    if transcript and transcript.segments:
        doc.add_page_break()
        doc.add_heading("Приложение. Расшифровка записи", level=1)
        has_time = any(s.end for s in transcript.segments)
        for s in transcript.segments:
            p = doc.add_paragraph()
            if has_time:
                from .transcribe import format_time

                run = p.add_run(f"[{format_time(s.start)}] ")
                run.font.size = Pt(9)
                run.bold = True
            p.add_run(s.text)

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
