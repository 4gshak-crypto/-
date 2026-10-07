import json
from pathlib import Path
from types import SimpleNamespace

import pytest

from meeting_notes.export import summary_to_markdown, write_docx
from meeting_notes.summarize import (
    ActionItem,
    MeetingSummary,
    SummaryError,
    Summarizer,
    Topic,
    build_user_message,
)
from meeting_notes.transcribe import Segment, Transcript


def sample_summary() -> MeetingSummary:
    return MeetingSummary(
        title="Планёрка по текущему ремонту",
        participants=["Иванов", "Филимонов"],
        summary="Обсудили ход работ по кровле и фасаду.",
        topics=[Topic(title="Кровля на Ленина, 5", discussion="Работы идут по графику, акт будет к пятнице.")],
        decisions=["Принять работы по кровле после акта."],
        action_items=[
            ActionItem(task="Подготовить акт по кровле", owner="Филимонов", due="до пятницы", context="«акт к пятнице сделаем»"),
            ActionItem(task="Уточнить смету по фасаду | подъезд 2", owner="не назначен", due="не назван", context=""),
        ],
        open_questions=["Кто оплачивает вывоз мусора?"],
        next_steps="Следующая планёрка в понедельник.",
    )


def sample_transcript() -> Transcript:
    return Transcript(
        segments=[Segment(0.0, 4.0, "Добрый день, начнём."), Segment(4.0, 9.0, "Акт по кровле к пятнице сделаем.")],
        language="ru", duration=9.0, source_name="planerka.mp3",
    )


def test_markdown_contains_sections_and_escapes_pipes():
    md = summary_to_markdown(sample_summary(), meeting_date="07.10.2026", source_name="planerka.mp3")
    assert md.startswith("# Планёрка по текущему ремонту")
    assert "**Дата:** 07.10.2026" in md
    for section in ("## Кратко", "## Обсуждали", "## Решения", "## Поручения", "## Открытые вопросы", "## Дальше"):
        assert section in md
    assert "| 1 | Подготовить акт по кровле | Филимонов | до пятницы |" in md
    assert "подъезд 2" in md and "| подъезд" not in md  # вертикальная черта внутри ячейки экранирована


def test_markdown_for_empty_summary():
    s = MeetingSummary(title="", participants=[], summary="", topics=[], decisions=[], action_items=[],
                       open_questions=[], next_steps="")
    md = summary_to_markdown(s)
    assert md.startswith("# Протокол совещания")
    assert md.count("—") >= 3


def test_docx_is_written(tmp_path: Path):
    out = tmp_path / "x" / "protocol.docx"
    write_docx(out, sample_summary(), sample_transcript(), meeting_date="07.10.2026")
    assert out.stat().st_size > 0
    from docx import Document

    doc = Document(str(out))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Планёрка по текущему ремонту" in text
    assert "Расшифровка записи" in text
    assert "[00:00:04] Акт по кровле к пятнице сделаем." in text
    assert doc.tables[0].rows[1].cells[2].text == "Филимонов"


def test_user_message_has_meta_and_timestamps():
    msg = build_user_message(sample_transcript(), "Планёрка", "07.10.2026", "Участники: Иванов")
    assert "Название совещания: Планёрка" in msg
    assert "Дата: 07.10.2026" in msg
    assert "Дополнительно от организатора: Участники: Иванов" in msg
    assert "<transcript>\n[00:00:00] Добрый день, начнём." in msg


class FakeStream:
    """Подмена client.beta.messages.stream: запоминает запрос и отдаёт готовый ответ."""

    def __init__(self, message, holder):
        self.message = message
        self.holder = holder

    def __call__(self, **request):
        self.holder.append(request)
        return self

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        return False

    def __iter__(self):
        yield SimpleNamespace(type="content_block_delta", delta=SimpleNamespace(type="text_delta", text="{}"))

    def get_final_message(self):
        return self.message


def make_client(message, holder):
    return SimpleNamespace(beta=SimpleNamespace(messages=SimpleNamespace(stream=FakeStream(message, holder))))


def test_summarizer_builds_request_and_returns_parsed():
    requests = []
    summary = sample_summary()
    message = SimpleNamespace(stop_reason="end_turn", parsed_output=summary)
    s = Summarizer(client=make_client(message, requests), model="claude-opus-5-5", effort="high", max_tokens=32000)
    progress = []
    result = s.summarize(sample_transcript(), meeting_title="Планёрка", on_progress=progress.append)
    assert result is summary
    req = requests[0]
    assert req["model"] == "claude-opus-5-5"
    assert req["output_format"] is MeetingSummary
    assert req["output_config"] == {"effort": "high"}
    assert req["fallbacks"] == "default"
    assert req["betas"] == ["server-side-fallback-2026-07-01"]
    assert "thinking" not in req  # Claude Opus 5.5: мышление включено всегда, параметр не передаётся
    assert req["messages"][0]["role"] == "user"
    assert "Планёрка" in req["messages"][0]["content"]
    assert progress and progress[0].startswith("Claude")


def test_summarizer_refusal_and_truncation():
    message = SimpleNamespace(stop_reason="refusal", parsed_output=None,
                              stop_details=SimpleNamespace(category="general_harms"))
    s = Summarizer(client=make_client(message, []), model="m")
    with pytest.raises(SummaryError, match="отказался.*general_harms"):
        s.summarize(sample_transcript())

    message = SimpleNamespace(stop_reason="max_tokens", parsed_output=None)
    s = Summarizer(client=make_client(message, []), model="m", max_tokens=100)
    with pytest.raises(SummaryError, match="100 токенов"):
        s.summarize(sample_transcript())


def test_summary_schema_is_acceptable_json_schema():
    schema = MeetingSummary.model_json_schema()
    assert set(schema["required"]) == {"title", "participants", "summary", "topics", "decisions",
                                       "action_items", "open_questions", "next_steps"}
    json.dumps(schema)
